#!/usr/bin/env python3
"""Export all OK ventes (STATUS=1 / STATUS_STATUS=1), colored, split at 12 months."""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from engine.processor import (  # noqa: E402
    DEFAULT_COLOR_FILLS,
    _assign_color,
    _canonical_tel_digits,
    _parse_date_column,
)
from engine.storage import (  # noqa: E402
    load_client_db,
    load_history,
    refresh_history_compact,
)

COLOR_ORDER = ["Green", "Blue", "Orange", "Red", "Unknown"]
COLOR_FR = {
    "Green": "Vert",
    "Blue": "Bleu",
    "Orange": "Orange",
    "Red": "Rouge",
    "Unknown": "Inconnu",
}
DISPLAY_COLS = [
    "TEL",
    "NOM",
    "PRENOM",
    "Date_Vente",
    "Date_Premiere_Vente",
    "Nb_Ventes",
    "Date_Dernier_Appel",
    "Jours_Depuis_Dernier_Appel",
    "Couleur",
    "Anciennete",
    "Statut_Actuel",
    "Lib_Status",
    "Agent_TV",
    "Duree_Dernier_Appel",
]


def _canon(tel: object) -> str:
    digits = "".join(ch for ch in str(tel or "") if ch.isdigit())
    return _canonical_tel_digits(digits)


def _status_code(val: object) -> str:
    try:
        return str(int(float(str(val).strip())))
    except (TypeError, ValueError):
        return str(val or "").replace(".0", "").strip()


def build_ventes_frame(hist: pd.DataFrame, db: pd.DataFrame) -> pd.DataFrame:
    work = hist.copy()
    work["TEL_K"] = work["TEL"].map(_canon)
    work = work[work["TEL_K"].astype(str).str.len() >= 9].copy()
    work["STATUS_N"] = work["STATUS"].map(_status_code)
    work["DATE_P"] = _parse_date_column(work["DATE"])

    ventes = work[work["STATUS_N"] == "1"].copy()
    if ventes.empty:
        return pd.DataFrame(columns=DISPLAY_COLS)

    sale_agg = ventes.groupby("TEL_K", as_index=False).agg(
        Date_Vente=("DATE_P", "max"),
        Date_Premiere_Vente=("DATE_P", "min"),
        Nb_Ventes=("DATE_P", "count"),
        TEL=("TEL", "last"),
        Lib_Status=("LIB_STATUS", "last"),
    )

    last = (
        work.sort_values("DATE_P")
        .groupby("TEL_K", as_index=False)
        .tail(1)[
            ["TEL_K", "DATE_P", "STATUS", "LIB_STATUS", "TV", "DUREE"]
        ]
        .rename(
            columns={
                "DATE_P": "Date_Dernier_Appel",
                "STATUS": "Statut_Actuel",
                "LIB_STATUS": "Lib_Dernier_Appel",
                "TV": "Agent_TV",
                "DUREE": "Duree_Dernier_Appel",
            }
        )
    )

    out = sale_agg.merge(last, on="TEL_K", how="left")
    today = pd.Timestamp(datetime.now().date())
    cutoff = today - pd.DateOffset(months=12)
    out["Jours_Depuis_Dernier_Appel"] = (today - out["Date_Dernier_Appel"]).dt.days
    out["Color"] = out["Jours_Depuis_Dernier_Appel"].map(_assign_color)
    out["Couleur"] = out["Color"].map(lambda c: COLOR_FR.get(str(c), str(c)))
    out["Anciennete"] = out["Date_Vente"].map(
        lambda d: (
            "Plus de 12 mois"
            if pd.notna(d) and d < cutoff
            else "Moins de 12 mois"
        )
    )
    lib = out["Lib_Status"].fillna("").astype(str).str.strip()
    out["Lib_Status"] = lib.where(lib != "", out.get("Lib_Dernier_Appel", "")).fillna("")

    db_work = db.copy()
    db_work["TEL_K"] = db_work["TEL"].map(_canon)
    name_cols = [c for c in ("NOM", "PRENOM", "Nom", "Prenom") if c in db_work.columns]
    keep = ["TEL_K"] + name_cols
    db_names = db_work[keep].drop_duplicates("TEL_K")
    if "Nom" in db_names.columns and "NOM" not in db_names.columns:
        db_names = db_names.rename(columns={"Nom": "NOM"})
    if "Prenom" in db_names.columns and "PRENOM" not in db_names.columns:
        db_names = db_names.rename(columns={"Prenom": "PRENOM"})
    out = out.merge(db_names, on="TEL_K", how="left")
    if "NOM" not in out.columns:
        out["NOM"] = ""
    if "PRENOM" not in out.columns:
        out["PRENOM"] = ""

    out["Color_Ordre"] = out["Color"].map(
        {name: i for i, name in enumerate(COLOR_ORDER)}
    ).fillna(99)
    out = out.sort_values(
        ["Color_Ordre", "Jours_Depuis_Dernier_Appel", "Date_Vente"],
        ascending=[True, False, False],
    ).reset_index(drop=True)

    for col in ("Date_Vente", "Date_Premiere_Vente", "Date_Dernier_Appel"):
        out[col] = pd.to_datetime(out[col], errors="coerce").dt.strftime("%d/%m/%Y")
    out["Jours_Depuis_Dernier_Appel"] = pd.to_numeric(
        out["Jours_Depuis_Dernier_Appel"], errors="coerce"
    ).fillna(0).astype(int)
    out["Nb_Ventes"] = pd.to_numeric(out["Nb_Ventes"], errors="coerce").fillna(0).astype(int)
    out["Statut_Actuel"] = out["Statut_Actuel"].map(_status_code)
    return out


def _style_sheet(ws, df: pd.DataFrame, *, header_row: int) -> None:
    fills = {
        "Vert": PatternFill(
            start_color=DEFAULT_COLOR_FILLS["Green"],
            end_color=DEFAULT_COLOR_FILLS["Green"],
            fill_type="solid",
        ),
        "Bleu": PatternFill(
            start_color=DEFAULT_COLOR_FILLS["Blue"],
            end_color=DEFAULT_COLOR_FILLS["Blue"],
            fill_type="solid",
        ),
        "Orange": PatternFill(
            start_color=DEFAULT_COLOR_FILLS["Orange"],
            end_color=DEFAULT_COLOR_FILLS["Orange"],
            fill_type="solid",
        ),
        "Rouge": PatternFill(
            start_color=DEFAULT_COLOR_FILLS["Red"],
            end_color=DEFAULT_COLOR_FILLS["Red"],
            fill_type="solid",
        ),
    }
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="00234E", end_color="00234E", fill_type="solid")
    for cell in ws[header_row]:
        cell.font = header_font
        cell.fill = header_fill
    headers = [c.value for c in ws[header_row]]
    if "Couleur" not in headers:
        return
    color_col = headers.index("Couleur") + 1
    for row in range(header_row + 1, header_row + 1 + len(df)):
        cell = ws.cell(row, color_col)
        fill = fills.get(str(cell.value))
        if fill:
            cell.fill = fill
    for idx, col in enumerate(df.columns, start=1):
        width = min(max([len(str(col))] + [len(str(v)) for v in df[col].head(80).tolist()]) + 2, 36)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = f"A{header_row + 1}"


def write_excel(df: pd.DataFrame, path: Path) -> None:
    plus = df[df["Anciennete"] == "Plus de 12 mois"].copy()
    moins = df[df["Anciennete"] == "Moins de 12 mois"].copy()
    by_color = (
        df.groupby("Couleur", dropna=False)
        .size()
        .reindex(["Vert", "Bleu", "Orange", "Rouge", "Inconnu"])
        .fillna(0)
        .astype(int)
    )
    summary = pd.DataFrame(
        {
            "Metric": [
                "TEL uniques OK vente",
                "Moins de 12 mois",
                "Plus de 12 mois",
                "Vert (>60 j depuis dernier appel)",
                "Bleu (39-60 j)",
                "Orange (15-38 j)",
                "Rouge (<15 j)",
                "Export",
            ],
            "Value": [
                len(df),
                len(moins),
                len(plus),
                int(by_color.get("Vert", 0)),
                int(by_color.get("Bleu", 0)),
                int(by_color.get("Orange", 0)),
                int(by_color.get("Rouge", 0)),
                datetime.now().strftime("%d/%m/%Y %H:%M"),
            ],
        }
    )
    export_all = df[DISPLAY_COLS]
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Resume", index=False)
        export_all.to_excel(writer, sheet_name="Toutes", index=False)
        moins[DISPLAY_COLS].to_excel(writer, sheet_name="Moins_12_mois", index=False)
        plus[DISPLAY_COLS].to_excel(writer, sheet_name="Plus_12_mois", index=False)
        _style_sheet(writer.sheets["Toutes"], export_all, header_row=1)
        _style_sheet(writer.sheets["Moins_12_mois"], moins[DISPLAY_COLS], header_row=1)
        _style_sheet(writer.sheets["Plus_12_mois"], plus[DISPLAY_COLS], header_row=1)
        ws = writer.sheets["Resume"]
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 18
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="00234E", end_color="00234E", fill_type="solid")


def main() -> None:
    print("Rebuild cache compact (STATUS + STATUS_STATUS)…", flush=True)
    refresh_history_compact()
    hist = load_history()
    db = load_client_db()
    print(f"hist compact {len(hist):,} · client_db {len(db):,}", flush=True)
    df = build_ventes_frame(hist, db)
    print(
        f"OK vente {len(df):,} · moins 12 mois {int((df['Anciennete']=='Moins de 12 mois').sum()):,} "
        f"· plus 12 mois {int((df['Anciennete']=='Plus de 12 mois').sum()):,}",
        flush=True,
    )
    out = PROJECT_ROOT / "data" / "exports" / "OK_Ventes_Couleurs_12mois.xlsx"
    dl = Path.home() / "Downloads" / "OK_Ventes_Couleurs_12mois.xlsx"
    write_excel(df, out)
    write_excel(df, dl)
    print(f"Excel → {out}")
    print(f"Excel → {dl}")


if __name__ == "__main__":
    main()
