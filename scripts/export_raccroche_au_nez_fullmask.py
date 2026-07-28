#!/usr/bin/env python3
"""Export Data Client — RACCROCHE AU NEZ, toutes colonnes, masque couleur complet."""

from __future__ import annotations

import io
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from engine import database  # noqa: E402
from engine.processor import (  # noqa: E402
    ALL_COLORS,
    DEFAULT_COLOR_FILLS,
    _order_export_columns,
    process_data,
)

STATUS_LABEL = "RACCROCHE AU NEZ"
OUTPUT_DIR = PROJECT_ROOT / "data" / "exports"


def _filter_raccroche(df: pd.DataFrame) -> pd.DataFrame:
    if "LIB_STATUS" not in df.columns:
        raise ValueError("Colonne LIB_STATUS introuvable dans la base client.")
    lib = df["LIB_STATUS"].astype(str).str.strip().str.upper()
    return df[lib == STATUS_LABEL.upper()].copy()


def _apply_color_mask(wb, *, header_row: int = 6) -> None:
    fills = {
        name: PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
        for name, hex_color in DEFAULT_COLOR_FILLS.items()
    }
    for ws in wb.worksheets:
        headers = [c.value for c in ws[header_row]]
        if "Color" not in headers:
            continue
        color_col = headers.index("Color") + 1
        for row in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(row, color_col)
            fill = fills.get(cell.value)
            if fill:
                cell.fill = fill
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[column_letter].width = min(max_length + 2, 42)


def export_raccroche_fullmask(output_path: Path | None = None) -> tuple[Path, int, dict[str, int]]:
    db = database.load_db()
    hist = database.load_history()
    subset = _filter_raccroche(db)

    latest, stats = process_data(
        df_db=db,
        df_hist=hist,
        drop_unmapped_status=False,
        return_stats=True,
    )
    enrich_cols = [
        c
        for c in (
            "Color",
            "Days_Since_Last_Call",
            "DATE",
            "STATUS",
            "Status_Category",
            "DATETIME",
            "HEURE",
            "DUREE",
            "LIB_DETAIL",
        )
        if c in latest.columns
    ]
    if enrich_cols:
        latest_slice = latest[["TEL", *enrich_cols]].drop_duplicates(subset=["TEL"], keep="last")
        # Keep client-db columns; only add hist columns not already present.
        add_cols = [c for c in enrich_cols if c not in subset.columns]
        if add_cols:
            subset = subset.merge(
                latest_slice[["TEL", *add_cols]],
                on="TEL",
                how="left",
            )

    if "Days_Since_Last_Call" in subset.columns:
        subset = subset.sort_values(
            ["Days_Since_Last_Call", "TEL"],
            ascending=[False, True],
            na_position="last",
        )

    color_counts = (
        subset["Color"].value_counts(dropna=False).astype(int).to_dict()
        if "Color" in subset.columns
        else {}
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if output_path is None:
        stamp = datetime.now().strftime("%Y%m%d")
        output_path = OUTPUT_DIR / f"Recyclage_Raccroche_Au_Nez_FullMask_{stamp}.xlsx"

    summary = pd.DataFrame(
        {
            "Indicateur": [
                "Statut (LIB_STATUS)",
                "Lignes exportées",
                "TEL uniques",
                "Couleurs incluses (full mask)",
            ],
            "Valeur": [
                STATUS_LABEL,
                len(subset),
                subset["TEL"].nunique() if "TEL" in subset.columns else len(subset),
                ", ".join(ALL_COLORS),
            ],
        }
    )
    for color in ALL_COLORS:
        summary = pd.concat(
            [
                summary,
                pd.DataFrame(
                    {
                        "Indicateur": [f"  · {color}"],
                        "Valeur": [color_counts.get(color, 0)],
                    }
                ),
            ],
            ignore_index=True,
        )

    cols = _order_export_columns(subset.columns.tolist())
    export_df = subset[cols]

    buffer = io.BytesIO()
    sheet = "Raccroche au nez"
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name=sheet[:31], index=False, startrow=0)
        export_df.to_excel(writer, sheet_name=sheet[:31], index=False, startrow=5)

    buffer.seek(0)
    wb = load_workbook(buffer)
    _apply_color_mask(wb)
    wb.save(output_path)

    return output_path, len(subset), color_counts


def main() -> None:
    from engine.config_env import bootstrap_env

    bootstrap_env()
    path, n, colors = export_raccroche_fullmask()
    print(f"Export : {path}")
    print(f"Lignes : {n:,}")
    if colors:
        print("Répartition couleur :", colors)


if __name__ == "__main__":
    main()
